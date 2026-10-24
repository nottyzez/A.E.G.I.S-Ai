"""
prepare_logs.py — Sync & Convert hydroponic_log.xlsx to tidy JSON, JS & CSV for HydroSync Dashboards.

Reads & Synchronizes:
  - 4 Foam box tanks (Control, ขาด N, ขาด P, ขาด K) [23.3L, groundwater baseline, up to 22/10/2026 09:00]
  - 8 Cup replicates (แก้ว 1 ถึง แก้ว 8) [750ml, drinking water baseline, un-swapped & up to 22/10/2026 09:00]

Outputs:
  - Ai มหาปลัย/dataset_real/hydroponic_log.xlsx (updated in-place with cleaned dates & 22/10/2026 09:00 rows)
  - data/sensor_log.json
  - data/sensor_data.js (window.SENSOR_DATA = ...)
  - server/static/data/sensor_data.js
  - data/google_sheets_export/*.csv (for easy Google Sheets import)
"""

import json
import os
import sys
import io
import csv
from datetime import datetime, date as ddate, time as dtime
from pathlib import Path

try:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")
except Exception:
    pass

import openpyxl

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent

EXCEL_PATH = PROJECT_ROOT / "Ai มหาปลัย" / "dataset_real" / "hydroponic_log.xlsx"
OUTPUT_JSON = SCRIPT_DIR / "sensor_log.json"
OUTPUT_JS = SCRIPT_DIR / "sensor_data.js"
STATIC_JS = PROJECT_ROOT / "server" / "static" / "data" / "sensor_data.js"
CSV_EXPORT_DIR = SCRIPT_DIR / "google_sheets_export"

TANK_MAP = {
    "Control": {"id": "Control", "name": "ถังโฟม Control", "group": "Control", "system": "tank_23L", "baseline_r1": 2.641, "baseline_r2": 3.000},
    "ขาด N":   {"id": "N",       "name": "ถังโฟม ขาด N",   "group": "N",       "system": "tank_23L", "baseline_r1": 2.572, "baseline_r2": 1.530},
    "ขาด P":   {"id": "P",       "name": "ถังโฟม ขาด P",   "group": "P",       "system": "tank_23L", "baseline_r1": 3.512, "baseline_r2": 1.912},
    "ขาด K":   {"id": "K",       "name": "ถังโฟม ขาด K",   "group": "K",       "system": "tank_23L", "baseline_r1": 2.196, "baseline_r2": 1.660},
}

CUP_MAP = {
    "แก้ว 1": {"id": "Cup1", "name": "แก้ว 1 (Control ซ้ำ 1)", "group": "Control", "system": "cup_750ml", "replicate": 1, "baseline_r1": 1.808},
    "แก้ว 2": {"id": "Cup2", "name": "แก้ว 2 (Control ซ้ำ 2)", "group": "Control", "system": "cup_750ml", "replicate": 2, "baseline_r1": 1.996},
    "แก้ว 3": {"id": "Cup3", "name": "แก้ว 3 (ขาด N ซ้ำ 1)",   "group": "N",       "system": "cup_750ml", "replicate": 1, "baseline_r1": 1.130},
    "แก้ว 4": {"id": "Cup4", "name": "แก้ว 4 (ขาด N ซ้ำ 2)",   "group": "N",       "system": "cup_750ml", "replicate": 2, "baseline_r1": 1.144},
    "แก้ว 5": {"id": "Cup5", "name": "แก้ว 5 (ขาด P ซ้ำ 1)",   "group": "P",       "system": "cup_750ml", "replicate": 1, "baseline_r1": 1.086},
    "แก้ว 6": {"id": "Cup6", "name": "แก้ว 6 (ขาด P ซ้ำ 2)",   "group": "P",       "system": "cup_750ml", "replicate": 2, "baseline_r1": 1.150},
    "แก้ว 7": {"id": "Cup7", "name": "แก้ว 7 (ขาด K ซ้ำ 1)",   "group": "K",       "system": "cup_750ml", "replicate": 1, "baseline_r1": 0.938},
    "แก้ว 8": {"id": "Cup8", "name": "แก้ว 8 (ขาด K ซ้ำ 2)",   "group": "K",       "system": "cup_750ml", "replicate": 2, "baseline_r1": 1.020},
}

HEADER_ROW = 6
DATA_START_ROW = 7
BASELINE_CELL = (4, 2)  # Row 4, Column B

OPTIMAL = {
    "ec_min": 1.2, "ec_max": 2.5,
    "ph_min": 5.5, "ph_max": 6.5,
    "temp_min": 22.0, "temp_max": 28.0,
}

DAY1_TANK = datetime(2026, 10, 9)
DAY1_CUP = datetime(2026, 10, 20)
TRANSPLANT_SHOCK_DAYS = 3

# ---------------------------------------------------------------------------
# Authoritative Cleaned Dataset up to 22/10/2026 09:00:00
# Format per row: (date_str_iso, time_str, ec, ph_before, ph_after, temp, image_name, notes, round_str, formula_str)
# ---------------------------------------------------------------------------
AUTHORITATIVE_SHEETS = {
    "Control": [
        ("2026-10-09", "08:36:00", 2.641, None, 6.39, 28.5, None, None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-09", "20:36:00", 2.596, None, 6.46, 25.5, None, None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-10", "08:00:00", None,  None, None, None, "10_10_2026-time08_00.jpg", "ไม่ได้วัด", "รอบ 1", "ครบธาตุ"),
        ("2026-10-10", "16:16:00", 2.436, None, None, 30.0, "10_10_2026-time16_16.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-11", "10:00:00", 2.580, 7.15, 5.67, 28.0, "11_10_2026-time10_00.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-11", "17:00:00", 2.800, 8.09, 5.60, 27.0, "11_10_2026-time17_00.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-12", "08:30:00", 2.520, 7.52, 6.22, 28.5, "12_10_2026-time08_30.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-12", "16:30:00", 2.438, 6.85, 6.85, 30.0, "12_10_2026-time16_30.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-13", "07:00:00", 2.676, 7.11, 5.97, 26.5, "13_10_2026-time07_00.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-13", "17:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "ครบธาตุ"),
        ("2026-10-14", "07:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "ครบธาตุ"),
        ("2026-10-14", "17:00:00", 2.648, 6.21, 6.21, 28.0, "14_10_2026-time17_00.jpg", None, "รอบ 1", "ครบธาตุ"),
        ("2026-10-15", "08:00:00", 3.000, 7.08, 6.45, None, "15_10_2026-time08_00.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-15", "17:30:00", 2.824, 6.86, 6.25, 28.5, "15_10_2026-time17_30.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-16", "08:30:00", 2.756, 7.82, 6.32, 28.5, "16_10_2026-time08_30.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-16", "16:30:00", 2.816, 7.96, 5.15, 30.0, "16_10_2026-time16_30.jpg", "เติมสารปรับ pH เยอะเกิน", "รอบ 2", "ครบธาตุ"),
        ("2026-10-17", "08:30:00", 2.680, 6.80, 5.66, 28.0, "17_10_2026-time08_30.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-17", "17:00:00", 2.868, 5.62, 5.62, 30.0, "17_10_2026-time17_00.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-18", "07:30:00", 3.102, 5.80, 5.80, 27.1, "18_10_2026-time07_30.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-18", "17:00:00", 2.908, 5.77, 5.77, 31.5, "18_10_2026-time17_00.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-19", "08:30:00", 3.044, 5.79, 5.79, 28.0, "19_10_2026-time08_30.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-19", "17:00:00", 2.888, 5.86, 5.86, 32.0, "19_10_2026-time17_00.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-20", "08:30:00", 3.002, 6.06, 6.06, 28.5, "20_10_2026-time08_30.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-20", "17:00:00", 2.958, 6.08, 6.08, 30.6, "20_10_2026-time17_00.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-21", "08:00:00", 3.006, 6.12, 6.12, 27.3, "21_10_2026-time08_00.jpg", None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-21", "18:00:00", 3.048, 5.77, 5.77, 28.5, None, None, "รอบ 2", "ครบธาตุ"),
        ("2026-10-22", "09:00:00", 2.872, 6.20, 6.20, 29.5, None, None, "รอบ 2", "ครบธาตุ"),
    ],
    "ขาด N": [
        ("2026-10-08", "17:53:00", 2.572, None, 5.96, 30.0, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-09", "08:56:00", 1.956, None, 6.43, 28.5, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-09", "20:42:00", 2.232, None, 6.25, 25.5, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-10", "08:00:00", None,  None, None, None, "10_10_2026-time08_00.jpg", "ลืมวัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-10", "16:18:00", 1.940, None, None, 30.0, "10_10_2026-time16_18.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-11", "10:00:00", 1.964, 7.14, 6.20, 28.0, "11_10_2026-time10_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-11", "17:00:00", 2.000, 6.36, 6.36, 27.0, "11_10_2026-time17_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-12", "08:30:00", 1.785, 6.35, 6.35, 28.5, "12_10_2026-time08_30.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-12", "16:30:00", 1.740, 6.39, 6.39, 30.0, "12_10_2026-time16_30.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-13", "07:00:00", 1.882, 6.34, 6.34, 26.5, "13_10_2026-time07_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-13", "17:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-14", "07:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-14", "17:00:00", 1.840, 6.32, 6.32, 28.0, "14_10_2026-time17_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-15", "08:00:00", 1.530, 7.00, 6.34, None, "15_10_2026-time08_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-15", "17:30:00", 1.826, 6.85, 6.45, 28.5, "15_10_2026-time17_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-16", "08:30:00", 1.770, 7.29, 6.45, None, "16_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-16", "16:30:00", 1.704, 7.05, 6.41, 30.0, "16_10_2026-time16_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-17", "08:30:00", 1.738, 6.85, 6.45, 28.0, "17_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-17", "17:00:00", 1.672, 6.47, 6.47, 30.0, "17_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-18", "07:30:00", 1.796, 6.62, 6.23, 27.1, "18_10_2026-time07_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-18", "17:00:00", 1.668, 6.45, 6.45, 31.5, "18_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-19", "08:30:00", 1.752, 6.56, 6.23, 29.5, "19_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-19", "17:00:00", 1.696, 6.33, 6.33, 30.6, "19_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-20", "08:30:00", 1.760, 6.21, 6.21, 28.5, "20_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-20", "17:00:00", 1.670, 6.25, 6.25, 30.6, "20_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-21", "08:00:00", 1.734, 6.33, 6.33, 27.3, "21_10_2026-time08_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-21", "18:00:00", 1.760, 6.23, 6.23, 28.5, None, None, "รอบ 2", "สูตร 25%"),
        ("2026-10-22", "09:00:00", 1.744, 6.08, 6.08, 29.5, None, None, "รอบ 2", "สูตร 25%"),
    ],
    "ขาด P": [
        ("2026-10-08", "18:00:00", 3.512, None, 6.12, 30.0, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-09", "08:58:00", 2.856, None, 6.31, 28.5, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-09", "20:49:00", 2.784, None, 6.40, 25.5, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-10", "08:00:00", None,  None, None, None, "10_10_2026-time08_00.jpg", "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-10", "16:20:00", 2.562, None, None, 30.0, "10_10_2026-time16_20.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-11", "10:00:00", 2.660, 5.89, 5.89, 28.0, "11_10_2026-time10_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-11", "17:00:00", 2.700, 5.92, 5.92, 27.0, "11_10_2026-time17_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-12", "08:30:00", 2.304, 5.65, 5.65, 28.5, "12_10_2026-time08_30.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-12", "16:30:00", 2.330, 5.67, 5.67, 30.0, "12_10_2026-time16_30.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-13", "07:00:00", 2.452, 5.54, 5.54, 26.5, "13_10_2026-time07_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-13", "17:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-14", "07:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-14", "17:00:00", 2.496, 5.54, 5.54, 28.0, "14_10_2026-time17_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-15", "08:00:00", 1.912, 7.36, 6.31, 30.0, "15_10_2026-time08_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-15", "17:30:00", 1.898, 7.46, 6.34, 28.5, "15_10_2026-time17_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-16", "08:30:00", 1.948, 8.62, 6.09, None, "16_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-16", "16:30:00", 1.940, 8.22, 5.64, 30.0, "16_10_2026-time16_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-17", "08:30:00", 1.856, 6.86, 5.75, 28.0, "17_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-17", "17:00:00", 1.868, 7.04, 6.18, 30.0, "17_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-18", "07:30:00", 2.000, 6.56, 6.11, 27.6, "18_10_2026-time07_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-18", "17:00:00", 1.958, 6.53, 6.53, 31.5, "18_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-19", "08:30:00", 2.038, 6.66, 5.11, 28.0, "19_10_2026-time08_30.jpg", "เติมสารปรับ pH ลงไป 2 ml", "รอบ 2", "สูตร 25%"),
        ("2026-10-19", "17:00:00", 1.962, 6.30, 6.30, 30.0, "19_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-20", "08:30:00", 2.026, 6.41, 6.41, 27.6, "20_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-20", "17:00:00", 1.920, 6.56, 6.32, 30.6, "20_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-21", "08:00:00", 2.070, 6.77, 5.43, 27.3, "21_10_2026-time08_00.jpg", "เติมสารปรับ pH ลงไป 1 ml", "รอบ 2", "สูตร 25%"),
        ("2026-10-21", "18:00:00", 2.028, 5.88, 5.88, 28.0, None, None, "รอบ 2", "สูตร 25%"),
        ("2026-10-22", "09:00:00", 1.948, 6.41, 6.41, 29.5, None, None, "รอบ 2", "สูตร 25%"),
    ],
    "ขาด K": [
        ("2026-10-08", "18:15:00", 2.196, None, 6.28, 30.0, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-09", "09:02:00", 1.786, None, 6.41, 28.5, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-09", "20:53:00", 1.670, None, 6.36, None, None, None, "รอบ 1", "สูตร 25%"),
        ("2026-10-10", "08:00:00", None,  None, None, None, "10_10_2026-time08_00.jpg", "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-10", "16:22:00", 1.362, None, None, 30.0, "10_10_2026-time16_22.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-11", "10:00:00", 1.320, None, 6.45, 27.5, "11_10_2026-time10_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-11", "17:00:00", 1.380, 6.47, 6.47, 27.0, "11_10_2026-time17_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-12", "08:30:00", 1.276, 6.29, 6.29, 28.0, "12_10_2026-time08_30.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-12", "16:30:00", 1.254, 6.11, 6.11, 30.0, "12_10_2026-time16_30.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-13", "07:00:00", 1.282, 6.09, 6.09, 26.5, "13_10_2026-time07_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-13", "17:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-14", "07:00:00", None,  None, None, None, None, "ไม่ได้วัด", "รอบ 1", "สูตร 25%"),
        ("2026-10-14", "17:00:00", 1.280, 6.03, 6.03, 28.0, "14_10_2026-time17_00.jpg", None, "รอบ 1", "สูตร 25%"),
        ("2026-10-15", "08:00:00", 1.660, 6.58, 6.29, None, "15_10_2026-time08_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-15", "17:30:00", 1.642, 7.31, 6.28, 27.0, "15_10_2026-time17_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-16", "08:30:00", 1.364, 7.34, 6.07, 28.5, "16_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-16", "16:30:00", 1.300, 6.52, 5.36, 30.6, "16_10_2026-time16_30.jpg", "เติมสารปรับ pH ฟอสฟอริกเข้าไป 1 ml", "รอบ 2", "สูตร 25%"),
        ("2026-10-17", "08:30:00", 1.284, 5.78, 5.65, 28.0, "17_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-17", "17:00:00", 1.312, 5.71, 5.71, 30.0, "17_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-18", "07:30:00", 1.346, 5.58, 5.58, 27.6, "18_10_2026-time07_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-18", "17:00:00", 1.312, 5.58, 5.58, 30.0, "18_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-19", "08:30:00", 1.332, 5.56, 5.56, 28.0, "19_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-19", "17:00:00", 1.332, 5.56, 5.56, 29.6, "19_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-20", "08:30:00", 1.350, 5.63, 5.63, 27.6, "20_10_2026-time08_30.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-20", "17:00:00", 1.278, 5.90, 5.90, 30.6, "20_10_2026-time17_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-21", "08:00:00", 1.336, 5.59, 5.59, 27.3, "21_10_2026-time08_00.jpg", None, "รอบ 2", "สูตร 25%"),
        ("2026-10-21", "18:00:00", 1.342, 5.70, 5.70, 28.0, None, None, "รอบ 2", "สูตร 25%"),
        ("2026-10-22", "09:00:00", 1.308, 5.73, 5.73, 29.5, None, None, "รอบ 2", "สูตร 25%"),
    ],
    "แก้ว 1": [
        ("2026-10-20", "16:00:00", 1.808, 5.75, 5.75, None, "Cup1_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 1.806, 6.15, 6.15, None, "Cup1_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 1.946, 6.21, 6.21, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 2.018, 6.43, 6.43, None, None, None, "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 2": [
        ("2026-10-20", "16:00:00", 1.996, 5.58, 5.58, None, "Cup2_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 1.920, 6.12, 6.12, None, "Cup2_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 2.078, 6.15, 6.15, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 1.956, 6.30, 6.30, None, None, None, "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 3": [
        ("2026-10-20", "16:00:00", 1.130, 5.70, 5.70, None, "Cup3_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 1.160, 6.56, 6.35, None, "Cup3_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 1.228, 6.48, 6.48, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 1.456, 5.82, 5.82, None, None, "หยดสารปรับ pH ฟอสฟอริก 1 หยด (0.05cc) ค่า pH ดิ่งลง 1.58 เปลี่ยนสารละลายใหม่", "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 4": [
        ("2026-10-20", "16:00:00", 1.144, 5.80, 5.80, None, "Cup4_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 1.184, 6.56, 6.25, None, "Cup4_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 1.208, 6.44, 6.44, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 0.916, 6.56, 6.56, None, None, None, "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 5": [
        ("2026-10-20", "16:00:00", 1.086, 6.51, 6.51, None, "Cup5_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 1.092, 7.20, 6.11, None, "Cup5_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 1.116, 6.23, 6.23, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 1.108, 6.14, 6.14, None, None, None, "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 6": [
        ("2026-10-20", "16:00:00", 1.150, 6.18, 6.18, None, "Cup6_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 1.130, 7.08, 6.43, None, "Cup6_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 1.182, 6.48, 6.48, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 1.132, 6.41, 6.41, None, None, None, "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 7": [
        ("2026-10-20", "16:00:00", 0.938, 7.44, 6.15, None, "Cup7_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 0.910, 6.86, 6.40, None, "Cup7_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 0.946, None, None, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 1.226, None, None, None, None, "หยดสารปรับ pH ฟอสฟอริก 1 หยด (0.05cc) ค่า pH ดิ่งลง 1.58 เปลี่ยนสารละลายใหม่", "ชุดแก้ว 750ml", None),
    ],
    "แก้ว 8": [
        ("2026-10-20", "16:00:00", 1.020, 6.34, 6.34, None, "Cup8_20261020_1600.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "08:00:00", 0.990, 6.97, 6.41, None, "Cup8_20261021_0800.jpg", None, "ชุดแก้ว 750ml", None),
        ("2026-10-21", "18:00:00", 1.032, 6.62, 6.62, None, None, None, "ชุดแก้ว 750ml", None),
        ("2026-10-22", "09:00:00", 1.196, 6.36, 6.36, None, None, "หยดสารปรับ pH ฟอสฟอริก 1 หยด (0.05cc) ค่า pH ดิ่งลง 1.58 เปลี่ยนสารละลายใหม่", "ชุดแก้ว 750ml", None),
    ],
}


def compute_ec_change_ratio(sheet_name, row_tuple):
    """
    Compute fractional EC change ((EC - ref_baseline) / ref_baseline):
      - For rows with 'เปลี่ยนสารละลายใหม่' in notes -> 0.0 (new solution baseline)
      - For Tank 'รอบ 2' from 15/10/2026 onwards -> uses baseline_r2 for N, P, K
        and for Control from 17/10/2026 17:00 onwards (matching user spreadsheet)
      - Otherwise -> uses baseline_r1
    """
    d_str, t_str, ec, _, _, _, _, notes, round_str, _ = row_tuple
    if ec is None:
        return None
    if notes and "เปลี่ยนสารละลายใหม่" in str(notes):
        return 0.0

    if sheet_name in TANK_MAP:
        meta = TANK_MAP[sheet_name]
        if sheet_name == "Control":
            # In user's sheet, Control switched reference to Baseline รอบ 2 (3.0) starting 17/10 17:00
            if (d_str > "2026-10-17") or (d_str == "2026-10-17" and t_str >= "17:00:00"):
                ref = meta["baseline_r2"]
            else:
                ref = meta["baseline_r1"]
        else:
            ref = meta["baseline_r2"] if round_str == "รอบ 2" else meta["baseline_r1"]
    else:
        ref = CUP_MAP[sheet_name]["baseline_r1"]

    return (ec - ref) / ref if ref else 0.0


def sync_excel_workbook():
    """Update hydroponic_log.xlsx in-place with the cleaned 22/10/2026 09:00 dataset."""
    if not EXCEL_PATH.exists():
        return
    wb = openpyxl.load_workbook(str(EXCEL_PATH))
    for sheet_name, rows in AUTHORITATIVE_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        # Update Baseline EC cell (Row 4, Col 2)
        if sheet_name in TANK_MAP:
            ws.cell(4, 2).value = TANK_MAP[sheet_name]["baseline_r1"]
            ws.cell(4, 4).value = TANK_MAP[sheet_name]["baseline_r2"]
        elif sheet_name in CUP_MAP:
            ws.cell(4, 2).value = CUP_MAP[sheet_name]["baseline_r1"]

        # Clear old data rows from DATA_START_ROW downwards
        max_r = max(ws.max_row, DATA_START_ROW + len(rows) + 5)
        for r_idx in range(DATA_START_ROW, max_r + 1):
            for c_idx in range(1, 12):
                ws.cell(r_idx, c_idx).value = None

        # Write authoritative rows
        for idx, r_tup in enumerate(rows):
            r_idx = DATA_START_ROW + idx
            d_str, t_str, ec, ph_b, ph_a, temp, img_name, notes, round_str, form_str = r_tup
            ec_ratio = compute_ec_change_ratio(sheet_name, r_tup)

            ws.cell(r_idx, 1).value = datetime.strptime(d_str, "%Y-%m-%d")
            ws.cell(r_idx, 1).number_format = "yyyy-mm-dd"
            ws.cell(r_idx, 2).value = t_str
            ws.cell(r_idx, 3).value = ec
            ws.cell(r_idx, 4).value = ph_b
            ws.cell(r_idx, 5).value = ph_a
            ws.cell(r_idx, 6).value = temp
            ws.cell(r_idx, 7).value = round(ec_ratio, 6) if ec_ratio is not None else None
            ws.cell(r_idx, 8).value = img_name
            ws.cell(r_idx, 9).value = notes
            ws.cell(r_idx, 10).value = round_str
            if form_str is not None:
                ws.cell(r_idx, 11).value = form_str

    wb.save(str(EXCEL_PATH))
    print(f"✅ Synchronized {EXCEL_PATH} with cleaned data up to 2026-10-22 09:00:00")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def parse_numeric(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if s in ("", "-", "-.", "None"):
        return None
    if s.endswith("%"):
        try:
            return float(s[:-1]) / 100.0
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_date_cell(val):
    """Parse date cell (datetime, date, YYYY-MM-DD, DD/MM/YYYY) and fix swapped Oct month/day."""
    if val is None:
        return None
    if isinstance(val, datetime):
        d = val.date()
    elif isinstance(val, ddate):
        d = val
    else:
        s = str(val).strip().split(" ")[0]
        d = None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                d = datetime.strptime(s, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            return None

    # Fix swapped day/month in October 2026 (e.g., 2026-09-10 -> 2026-10-09, 2026-11-10 -> 2026-10-11, 2026-12-10 -> 2026-10-12)
    if d.year == 2026 and d.day == 10 and d.month in (9, 11, 12):
        d = ddate(2026, 10, d.month)
    return d


def parse_time_cell(val):
    if val is None:
        return None
    if isinstance(val, dtime):
        return val
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, (int, float)):
        h = int(val)
        m_frac = val - h
        m = int(round(m_frac * 100))
        if m >= 60:
            m = int(round(m_frac * 60))
        return dtime(h, m)
    s = str(val).strip()
    for fmt in ("%H:%M", "%H:%M:%S", "%H.%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    return None


def make_ts(date_val, time_val):
    d = parse_date_cell(date_val)
    if d is None:
        return None
    if time_val is None:
        return datetime.combine(d, dtime(0, 0)).isoformat()
    return datetime.combine(d, time_val).isoformat()


def detect_quality_flags(unit_key, system_type, ts_str, ec, ph_before, ph_after, notes, baseline_ec):
    flags = []
    day1_ref = DAY1_CUP if system_type == "cup_750ml" else DAY1_TANK
    if ts_str:
        try:
            ts_dt = datetime.fromisoformat(ts_str)
            day_num = (ts_dt.date() - day1_ref.date()).days + 1
            if 1 <= day_num <= TRANSPLANT_SHOCK_DAYS:
                flags.append("transplant_shock")
        except (ValueError, TypeError):
            pass

    if ec is None:
        flags.append("missing_ec")
    if ph_after is None and ph_before is None:
        flags.append("missing_ph")

    if ph_before is not None and ph_after is not None:
        if abs(ph_before - ph_after) > 1.5:
            flags.append("ph_anomaly")

    if unit_key == "P" and baseline_ec and baseline_ec > 3.0:
        flags.append("high_baseline")

    if notes:
        notes_str = str(notes)
        notes_lower = notes_str.lower()
        if "เปลี่ยนสารละลายใหม่" in notes_str:
            flags.append("solution_replaced")
        if "1.58" in notes_str or ("เยอะ" in notes_lower or "ผิด" in notes_lower or "เกิน" in notes_lower or "ดิ่ง" in notes_lower):
            flags.append("acid_error")
        if "เติม" in notes_str or "หยด" in notes_str:
            flags.append("acid_addition")
        if "ลืมวัด" in notes_str or "ไม่ได้วัด" in notes_str or "ไม่่ได้วัด" in notes_str:
            flags.append("missed_measurement")

    return flags if flags else None


def export_sheets_to_csv(wb, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        filename = f"{sheet_name}.csv".replace(" ", "_")
        csv_path = output_dir / filename
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    writer.writerow(["" if c is None else str(c) for c in row])
    print(f"📁 Exported all {len(wb.sheetnames)} sheets to CSV in: {output_dir}")


def main():
    sync_excel_workbook()

    print(f"Reading {EXCEL_PATH} ...")
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    export_sheets_to_csv(wb, CSV_EXPORT_DIR)

    baselines = {}
    baselines_round2 = {}
    readings = []
    quality_events = []
    units_meta = {}

    all_sheets_to_process = []
    for sheet_name, meta in TANK_MAP.items():
        all_sheets_to_process.append((sheet_name, meta))
    for sheet_name, meta in CUP_MAP.items():
        all_sheets_to_process.append((sheet_name, meta))

    for sheet_name, meta in all_sheets_to_process:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        unit_id = meta["id"]
        system_type = meta["system"]
        group_name = meta["group"]

        baseline = parse_numeric(ws.cell(*BASELINE_CELL).value) or meta["baseline_r1"]
        baseline_r2 = meta.get("baseline_r2")
        baselines[unit_id] = baseline
        if baseline_r2:
            baselines_round2[unit_id] = baseline_r2

        units_meta[unit_id] = {
            "id": unit_id,
            "name": meta["name"],
            "sheet_name": sheet_name,
            "group": group_name,
            "system": system_type,
            "baseline_ec": baseline,
            "baseline_ec_round2": baseline_r2,
            "replicate": meta.get("replicate", 1),
        }

        row = DATA_START_ROW
        prev_ph_after = None

        while row <= ws.max_row:
            date_val = ws.cell(row, 1).value
            if date_val is None:
                row += 1
                continue

            time_val = parse_time_cell(ws.cell(row, 2).value)
            ec = parse_numeric(ws.cell(row, 3).value)
            ph_before = parse_numeric(ws.cell(row, 4).value)
            ph_after = parse_numeric(ws.cell(row, 5).value)
            temp = parse_numeric(ws.cell(row, 6).value)
            ec_change = parse_numeric(ws.cell(row, 7).value)
            image_name = ws.cell(row, 8).value
            notes = ws.cell(row, 9).value
            solution_round = ws.cell(row, 10).value

            ts_str = make_ts(date_val, time_val)

            if notes and "เปลี่ยนสารละลายใหม่" in str(notes):
                ec_change_pct = 0.0
            elif ec_change is not None and ec_change != -1.0:
                ec_change_pct = round(ec_change * 100, 2)
            elif ec is not None and baseline:
                ref_b = baseline_r2 if (solution_round == "รอบ 2" and baseline_r2) else baseline
                ec_change_pct = round(((ec - ref_b) / ref_b) * 100, 2)
            else:
                ec_change_pct = None

            flags = detect_quality_flags(
                unit_id, system_type, ts_str, ec, ph_before, ph_after, notes, baseline
            )

            if prev_ph_after is not None and ph_before is not None:
                if abs(ph_before - prev_ph_after) > 1.5:
                    if flags is None:
                        flags = []
                    flags.append("ph_jump_between_readings")
                    quality_events.append({
                        "tank": unit_id,
                        "system": system_type,
                        "ts": ts_str,
                        "type": "ph_jump",
                        "detail": f"pH jumped from {prev_ph_after} to {ph_before}",
                    })

            if notes and "เปลี่ยนสารละลายใหม่" in str(notes):
                quality_events.append({
                    "tank": unit_id,
                    "system": system_type,
                    "ts": ts_str,
                    "type": "solution_replaced",
                    "detail": str(notes).strip(),
                })

            reading = {
                "tank": unit_id,
                "system": system_type,
                "group": group_name,
                "ts": ts_str,
                "ec": ec,
                "ph_before": ph_before,
                "ph_after": ph_after,
                "temp": temp,
                "ec_change_pct": ec_change_pct,
                "solution_round": str(solution_round).strip() if solution_round else None,
                "image": str(image_name).strip() if image_name else None,
                "notes": str(notes).strip() if notes else None,
                "quality_flag": flags,
            }
            readings.append(reading)

            if ph_after is not None:
                prev_ph_after = ph_after

            row += 1

    seen = set()
    unique_events = []
    for ev in quality_events:
        key = (ev["tank"], ev["ts"], ev["type"])
        if key not in seen:
            seen.add(key)
            unique_events.append(ev)

    output = {
        "generated_at": datetime.now().isoformat(),
        "source": str(EXCEL_PATH.name),
        "optimal_ranges": OPTIMAL,
        "baselines": baselines,
        "baselines_round2": baselines_round2,
        "units": units_meta,
        "readings": readings,
        "quality_events": unique_events,
    }

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    js_content = "window.SENSOR_DATA = " + json.dumps(output, ensure_ascii=False) + ";\n"
    with open(OUTPUT_JS, "w", encoding="utf-8") as f:
        f.write(js_content)

    STATIC_JS.parent.mkdir(parents=True, exist_ok=True)
    with open(STATIC_JS, "w", encoding="utf-8") as f:
        f.write(js_content)

    print(f"\n✅ Output successfully written to:")
    print(f"   - {OUTPUT_JSON}")
    print(f"   - {OUTPUT_JS}")
    print(f"   - {STATIC_JS}")
    print(f"   Total units processed: {len(units_meta)} (4 Tanks + 8 Cups)")
    print(f"   Total readings: {len(readings)}")


if __name__ == "__main__":
    main()
